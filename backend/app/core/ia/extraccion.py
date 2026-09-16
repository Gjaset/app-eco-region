"""Extracción de texto plano de documentos subidos (PDF, DOCX, XLSX).

El texto extraído alimenta al validador IA; se trunca para controlar tokens.
"""
import io
import logging
import zipfile

logger = logging.getLogger(__name__)

MAX_CARACTERES = 20_000


class SinTextoError(ValueError):
    """El archivo no tiene texto extraíble (típicamente un escaneo)."""


def _extraer_pdf(contenido: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(contenido))
    partes = [pagina.extract_text() or "" for pagina in reader.pages]
    return "\n".join(partes)


def _extraer_docx(contenido: bytes) -> str:
    from docx import Document

    documento = Document(io.BytesIO(contenido))
    partes = [p.text for p in documento.paragraphs if p.text.strip()]
    # Tablas: la mayoría de campos de los formatos viven en tablas.
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells]
            partes.append(" | ".join(c for c in celdas if c))
    return "\n".join(partes)


def _extraer_xlsx(contenido: bytes) -> str:
    from openpyxl import load_workbook

    libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    partes: list[str] = []
    for hoja in libro.worksheets:
        partes.append(f"=== Hoja: {hoja.title} ===")
        for fila in hoja.iter_rows(values_only=True):
            valores = [str(c).strip() for c in fila if c is not None and str(c).strip()]
            if valores:
                partes.append(" | ".join(valores))
    libro.close()
    return "\n".join(partes)


def extraer_texto(nombre_archivo: str, contenido: bytes) -> tuple[str, bool]:
    """Devuelve (texto, truncado). Lanza SinTextoError si no hay texto útil."""
    nombre = nombre_archivo.lower()
    if not contenido:
        raise SinTextoError("El archivo está vacío.")

    def _seguro(fn):
        try:
            return fn(contenido)
        except SinTextoError:
            raise
        except Exception as exc:
            raise SinTextoError(f"No se pudo leer {nombre_archivo}: archivo inválido o corrupto.") from exc

    if nombre.endswith(".pdf"):
        texto = _seguro(_extraer_pdf)
    elif nombre.endswith(".docx"):
        texto = _seguro(_extraer_docx)
    elif nombre.endswith(".xlsx"):
        texto = _seguro(_extraer_xlsx)
    elif nombre.endswith(".xls"):
        # Formato binario antiguo: sin soporte directo. Usamos heurística mínima
        # via zip (no suele aplicar) y reportamos sin texto.
        if zipfile.is_zipfile(io.BytesIO(contenido)):
            texto = " ".join(
                zf.read(n).decode("utf-8", "ignore") for zf in [zipfile.ZipFile(io.BytesIO(contenido))] for n in zf.namelist()
            )
        else:
            raise SinTextoError("Los archivos .xls binarios no tienen texto extraíble; súbalo como .xlsx o PDF.")
    else:
        # Último recurso: intentar texto plano.
        try:
            texto = contenido.decode("utf-8", "ignore")
        except Exception as exc:  # pragma: no cover
            raise SinTextoError(f"No se pudo extraer texto de {nombre_archivo}") from exc

    texto = "\n".join(linea for linea in (p.strip() for p in texto.splitlines()) if linea)
    if len(texto.split()) < 10:
        raise SinTextoError(
            "El documento no contiene texto legible; si es un escaneo, se requiere revisión manual."
        )
    truncado = len(texto) > MAX_CARACTERES
    return texto[:MAX_CARACTERES], truncado

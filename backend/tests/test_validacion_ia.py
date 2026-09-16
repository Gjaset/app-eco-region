"""Pruebas de extracción de texto y normalización del dictamen IA."""
import io

import pytest

from app.core.ia.extraccion import SinTextoError, extraer_texto
from app.core.ia.validacion import _normalizar, dictamen_no_disponible


def _xlsx_con_datos() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Nombre del solicitante"
    ws["B1"] = "Juan Pérez Gómez"
    ws["A2"] = "Municipio"
    ws["B2"] = "Tunja"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _docx_con_tabla() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Ficha técnica de registro de individuos forestales")
    doc.add_paragraph("Radicado No. 2024-0001, fecha de visita técnica 12/05/2024")
    doc.add_paragraph("Solicitante: Juan Pérez Gómez, c.c. 1.032.456.789")
    tabla = doc.add_table(rows=1, cols=2)
    tabla.rows[0].cells[0].text = "Especie"
    tabla.rows[0].cells[1].text = "Eucalyptus globulus"
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def test_extraer_xlsx_incluye_celdas():
    texto, truncado = extraer_texto("f1.xlsx", _xlsx_con_datos())
    assert "Juan Pérez" in texto and "Tunja" in texto
    assert truncado is False


def test_extraer_docx_incluye_tablas():
    texto, _ = extraer_texto("f3.docx", _docx_con_tabla())
    assert "Eucalyptus globulus" in texto
    assert "|" in texto  # formato de tabla


def test_extraer_archivo_sin_texto_falla():
    with pytest.raises(SinTextoError):
        extraer_texto("escaneo.pdf", b"\x00\x01garbage-not-a-pdf")


def test_normalizar_repara_tipos():
    dictamen = _normalizar({"apto": True, "puntaje": "85", "observaciones": "ok"})
    assert dictamen == {
        "apto": True,
        "puntaje": 85,
        "faltantes": [],
        "observaciones": ["ok"],
        "resumen": "Documento sin faltantes apreciables.",
    }


def test_dictamen_no_disponible_no_lanza():
    d = dictamen_no_disponible("motivo X")
    assert d["apto"] is None and d["observaciones"] == ["motivo X"]


def test_parsear_json_con_comas_colgantes():
    from app.core.ia.validacion import _parse_json, _reparar_json

    roto = """```json
    {
      "apto": false,
      "puntaje": 45,
      "faltantes": [
        {"campo": "Especie", "detalle": "vacío"},
      ],
      "observaciones": ["Falta firma"],
      "resumen": "Texto con
salto de línea suelto"
    }
    ```
    """
    resultado = _parse_json(roto)
    assert resultado["apto"] is False
    assert resultado["puntaje"] == 45
    assert resultado["faltantes"][0]["campo"] == "Especie"


def test_reparar_json_basico():
    from app.core.ia.validacion import _reparar_json

    import json as _json
    assert _json.loads(_reparar_json('{"a": [1, 2,], }')) == {"a": [1, 2]}
